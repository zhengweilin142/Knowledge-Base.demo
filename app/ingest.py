"""文档解析 → 切分 → 生成可入库的 chunk。

设计目标：
1. 覆盖企业常见格式：pdf / docx / xlsx / txt / md；
2. 每个 chunk 都带来源定位（文件名、页码或章节），保证答案可溯源；
3. 切分按中文标点断句、带重叠，避免一句话被切断导致语义丢失。
"""
from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

from . import config

# --------------------------------------------------------------------------
# 数据结构
# --------------------------------------------------------------------------


@dataclass
class ParsedDoc:
    """解析结果：全文 + 用于溯源的位置段落。"""

    text: str
    segments: list[dict[str, Any]] = field(default_factory=list)

    def located_segments(self) -> list[dict[str, Any]]:
        if self.segments:
            return self.segments
        return [{"text": self.text, "page": None, "section": None}]


@dataclass
class Chunk:
    chunk_id: str
    text: str
    metadata: dict[str, Any]

    def as_chroma(self) -> dict[str, Any]:
        return {"id": self.chunk_id, "document": self.text, "metadata": self.metadata}


# --------------------------------------------------------------------------
# 解析器
# --------------------------------------------------------------------------

_SUPPORTED = {".pdf", ".docx", ".xlsx", ".xlsm", ".txt", ".md", ".markdown", ".csv"}


def supported_extensions() -> set[str]:
    return set(_SUPPORTED)


def _decode(raw: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "gb18030", "gbk", "big5"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _parse_text(path: Path) -> ParsedDoc:
    text = _decode(path.read_bytes())
    return ParsedDoc(text=text, segments=[{"text": text, "page": None, "section": None}])


def _parse_pdf(path: Path) -> ParsedDoc:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    segments: list[dict[str, Any]] = []
    for page_no, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception as exc:  # 单页解析失败不应中断整篇
            page_text = f"[本页解析失败: {exc}]"
        page_text = page_text.strip()
        if page_text:
            segments.append({"text": page_text, "page": page_no, "section": None})
    full = "\n\n".join(s["text"] for s in segments)
    return ParsedDoc(text=full, segments=segments)


def _parse_docx(path: Path) -> ParsedDoc:
    import docx  # python-docx

    document = docx.Document(str(path))
    segments: list[dict[str, Any]] = []
    current_section: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            segments.append(
                {"text": "\n".join(buffer).strip(), "page": None, "section": current_section}
            )
            buffer.clear()

    for para in document.paragraphs:
        content = (para.text or "").strip()
        if not content:
            continue
        style = (para.style.name or "") if para.style is not None else ""
        if style.startswith("Heading") or style in ("标题", "Title"):
            # 遇到标题：先把上一节收尾，标题本身也作为一节内容（利于召回小节标题）
            flush()
            current_section = content
            buffer.append(content)
        else:
            buffer.append(content)
    flush()

    for table_no, table in enumerate(document.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            cells = [(c.text or "").strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            segments.append(
                {
                    "text": f"[表格 {table_no}]\n" + "\n".join(rows),
                    "page": None,
                    "section": current_section,
                }
            )

    full = "\n\n".join(s["text"] for s in segments)
    return ParsedDoc(text=full, segments=segments)


def _parse_xlsx(path: Path) -> ParsedDoc:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    segments: list[dict[str, Any]] = []
    try:
        for ws in wb.worksheets:
            lines: list[str] = []
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    lines.append(" | ".join(cells))
                if len(lines) >= 300:  # 单表最多取 300 行，避免超大表格撑爆上下文
                    lines.append("…（本表行数过多已截断）")
                    break
            if lines:
                segments.append(
                    {"text": f"[工作表: {ws.title}]\n" + "\n".join(lines), "page": None, "section": ws.title}
                )
    finally:
        wb.close()
    full = "\n\n".join(s["text"] for s in segments)
    return ParsedDoc(text=full, segments=segments)


def parse_file(path: Path) -> ParsedDoc:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix in (".xlsx", ".xlsm"):
        return _parse_xlsx(path)
    if suffix in (".txt", ".md", ".markdown", ".csv"):
        return _parse_text(path)
    raise ValueError(f"暂不支持的格式: {suffix}（支持 {', '.join(sorted(_SUPPORTED))}）")


# --------------------------------------------------------------------------
# 切分
# --------------------------------------------------------------------------

# 先按段落切，再按中文断句切
_PARA_RE = re.compile(r"\n\s*\n+")
_SENT_RE = re.compile(r"(?<=[。！？；!?;])(?![”’）)])")


def _split_segments(text: str, target: int) -> list[str]:
    """把一段文本切成不超过 target 的小句单元。"""
    units: list[str] = []
    for para in _PARA_RE.split(text):
        para = para.strip()
        if not para:
            continue
        if len(para) <= target:
            units.append(para)
            continue
        sentences = [s.strip() for s in _SENT_RE.split(para) if s.strip()]
        for sentence in sentences:
            if len(sentence) <= target:
                units.append(sentence)
            else:  # 超长句（如无标点的表格行）硬切
                for i in range(0, len(sentence), target):
                    units.append(sentence[i : i + target])
    return units


def chunk_document(
    path: Path,
    doc: ParsedDoc,
    chunk_size: int | None = None,
    overlap: int | None = None,
) -> list[Chunk]:
    chunk_size = chunk_size or config.CHUNK_SIZE
    overlap = overlap or config.CHUNK_OVERLAP
    source_name = path.name
    chunks: list[Chunk] = []

    for seg_index, seg in enumerate(doc.located_segments()):
        seg_text = (seg.get("text") or "").strip()
        if not seg_text:
            continue
        units = _split_segments(seg_text, chunk_size)
        current: list[str] = []
        current_len = 0

        def emit(parts: list[str], first_of_segment: bool = False) -> None:
            body = "\n".join(parts).strip()
            # 阈值用于过滤页眉页脚之类的碎片；但每份文档至少保留一块，
            # 否则短文档（如摘要、便签）会被静默丢弃。
            if len(body) < config.MIN_CHUNK_SIZE and not first_of_segment:
                return
            digest = hashlib.sha1(f"{source_name}|{seg_index}|{len(chunks)}|{body[:200]}".encode()).hexdigest()[:16]
            chunks.append(
                Chunk(
                    chunk_id=f"{path.stem}-{digest}",
                    text=body,
                    metadata={
                        "source": source_name,
                        "file_path": str(path),
                        "ext": path.suffix.lower(),
                        "page": seg.get("page") if seg.get("page") is not None else -1,
                        "section": seg.get("section") or "",
                        "segment": seg_index,
                        "ordinal": len(chunks),
                    },
                )
            )

        for unit in units:
            if current and current_len + len(unit) + 1 > chunk_size:
                emit(current, first_of_segment=not chunks)
                # 构造重叠：从上一块尾部按句子回取，直到达到 overlap 长度
                tail: list[str] = []
                tail_len = 0
                for prev in reversed(current):
                    if tail_len >= overlap:
                        break
                    tail.insert(0, prev)
                    tail_len += len(prev) + 1
                current = tail
                current_len = tail_len
            current.append(unit)
            current_len += len(unit) + 1
        emit(current, first_of_segment=not chunks)

    return chunks


def iter_source_files(source_dir: Path | None = None) -> Iterable[Path]:
    source_dir = source_dir or config.SOURCE_DIR
    for path in sorted(source_dir.rglob("*")):
        if path.is_file() and path.suffix.lower() in _SUPPORTED:
            yield path


def prepare_chunks(path: Path) -> list[Chunk]:
    return chunk_document(path, parse_file(path))
