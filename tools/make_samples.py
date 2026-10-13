"""生成三类真实测试文档（Word / Excel / PDF），用于验证解析与切分链路。

PDF 优先用 LibreOffice 从中文 Word 转换（这样能测中文 PDF 提取）；
找不到 LibreOffice 时退化为最小可用 PDF，仅用于验证链路连通性。
"""
from __future__ import annotations

import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "source"
SRC.mkdir(parents=True, exist_ok=True)


def make_docx() -> Path:
    import docx

    doc = docx.Document()
    doc.add_heading("星海科技差旅管理制度", level=1)
    doc.add_heading("一、住宿标准", level=2)
    doc.add_paragraph("一线城市每晚不超过600元，二线城市不超过450元，其他城市不超过350元。")
    doc.add_paragraph("超标部分须事先取得部门总监书面批准，报销时附批准邮件截图。")
    doc.add_heading("二、交通标准", level=2)
    doc.add_paragraph("市内交通凭票实报；高铁乘坐二等座；飞机乘坐经济舱，飞行时间超过5小时可乘公务舱。")
    doc.add_heading("三、审批权限", level=2)
    table = doc.add_table(rows=3, cols=3)
    rows = [
        ("城市级别", "住宿上限", "审批人"),
        ("一线", "600元/晚", "部门总监"),
        ("二线", "450元/晚", "部门经理"),
    ]
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            table.rows[r].cells[c].text = text
    path = SRC / "差旅管理制度.docx"
    doc.save(str(path))
    return path


def make_xlsx() -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "供应商清单"
    ws.append(["供应商编号", "名称", "类别", "联系人", "电话"])
    ws.append(["SUP-001", "华信办公用品", "办公耗材", "王磊", "13800000001"])
    ws.append(["SUP-002", "远航网络科技", "IT设备", "李静", "13800000002"])
    ws.append(["SUP-003", "恒安安保服务", "安保", "张伟", "13800000003"])
    ws2 = wb.create_sheet("采购审批权限")
    ws2.append(["金额区间", "审批人"])
    ws2.append(["2000元以下", "部门负责人"])
    ws2.append(["2000-20000元", "分管副总"])
    ws2.append(["20000元以上", "总经理办公会"])
    path = SRC / "供应商与采购权限.xlsx"
    wb.save(str(path))
    return path


def find_libreoffice() -> str | None:
    for name in ("soffice", "soffice.exe", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    candidates = [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    for cand in candidates:
        if Path(cand).exists():
            return cand
    return None


def make_pdf_from(docx_path: Path) -> tuple[Path | None, str]:
    soffice = find_libreoffice()
    if not soffice:
        return None, "未找到 LibreOffice"
    out = SRC / f"{docx_path.stem}.pdf"
    try:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(SRC), str(docx_path)],
            check=True, capture_output=True, timeout=180,
        )
        if out.exists():
            return out, f"LibreOffice 转换成功 ({out.stat().st_size} 字节)"
        return None, "转换命令已执行但未生成 PDF"
    except Exception as exc:
        return None, f"转换失败: {exc}"


def make_minimal_pdf() -> Path:
    """兜底：最小合法 PDF（英文内容），仅验证解析链路。"""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
    ]
    stream = b"BT /F1 12 Tf 50 780 Td (Information Security Policy - Page 1) Tj ET"
    objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects)+1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode()

    path = SRC / "安全规范摘要.pdf"
    path.write_bytes(bytes(out))
    return path


def main() -> int:
    docx_path = make_docx()
    print(f"[生成] Word : {docx_path.name}")

    xlsx_path = make_xlsx()
    print(f"[生成] Excel: {xlsx_path.name}")

    pdf_path, note = make_pdf_from(docx_path)
    if pdf_path:
        print(f"[生成] PDF  : {pdf_path.name}  ({note})")
    else:
        fallback = make_minimal_pdf()
        print(f"[生成] PDF  : {fallback.name}  (兜底方案，原因：{note})")

    print("\n当前 data/source 内容：")
    for f in sorted(SRC.iterdir()):
        if f.is_file():
            print(f"  {f.stat().st_size:>9,} 字节  {f.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
